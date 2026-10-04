"""Excel versions of quotes and proforma invoices: the same layouts as the PDFs, with
formulas for line and section totals, red description lines and embedded photos."""

from __future__ import annotations

import io
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.worksheet import Worksheet

from app.config import AppSettings
from app.models import ProformaInvoice, Quote, QuoteSection
from app.services.documents.common import dmy, image_size, labels, ymd_slash
from app.services.documents.quote_pdf import ImageResolver, list_description
from app.services.quotes import printed_numbers

THIN = Side(style="thin", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
GREY = PatternFill("solid", fgColor="BFBFBF")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=True)
BOLD = Font(name="Calibri", bold=True, size=10)
NORMAL = Font(name="Calibri", size=10)
LINE_PT = 13.0  # row height per printed line
PHOTO_PX = 70

OFFER_WIDTHS = (5, 18, 55, 7, 10, 14)
LIST_WIDTHS = (5, 45, 20, 10, 26, 11, 11, 17)


def _money_format(value: Decimal | None, dollar: bool, always: bool) -> str:
    whole = value is not None and Decimal(value) == Decimal(value).to_integral_value()
    body = "#,##0" if whole and not always else "#,##0.00"
    return f'"$"{body}' if dollar else body


def _rich(lines: list[dict[str, Any]]) -> CellRichText | str:
    if not any(d.get("emphasis") for d in lines):
        return "\n".join(d.get("text", "") for d in lines)
    blocks: list[str | TextBlock] = []
    for index, d in enumerate(lines):
        text = d.get("text", "") + ("\n" if index < len(lines) - 1 else "")
        font = InlineFont(rFont="Calibri", sz=10, color="FFE60000" if d.get("emphasis") else None)
        blocks.append(TextBlock(font, text))
    return CellRichText(blocks)


def _set(ws: Worksheet, ref: str, value: Any, font: Font = NORMAL, align: Alignment = CENTER,
         fill: PatternFill | None = None, number_format: str | None = None,
         border: bool = True) -> None:  # fmt: skip
    cell = ws[ref]
    cell.value = value
    cell.font = font
    cell.alignment = align
    if border:
        cell.border = BOX
    if fill is not None:
        cell.fill = fill
    if number_format:
        cell.number_format = number_format


def _box_merge(ws: Worksheet, first: str, last: str) -> None:
    ws.merge_cells(f"{first}:{last}")
    for row in ws[f"{first}:{last}"]:
        for cell in row:
            cell.border = BOX


def _photos(ws: Worksheet, paths: list[Path], anchor_col: int, row: int) -> float:
    """Up to two photos side by side inside one cell; returns the row height they need."""
    x_px = 4
    for path in paths[:2]:
        w, h = image_size(path)
        img = XLImage(str(path))
        img.height = PHOTO_PX
        img.width = int(PHOTO_PX * w / h) if h else PHOTO_PX
        img.anchor = OneCellAnchor(
            _from=AnchorMarker(
                col=anchor_col - 1, colOff=pixels_to_EMU(x_px), row=row - 1, rowOff=pixels_to_EMU(6)
            ),
            ext=XDRPositiveSize2D(pixels_to_EMU(img.width), pixels_to_EMU(img.height)),
        )
        ws.add_image(img)
        x_px += img.width + 6
    return PHOTO_PX * 0.75 + 10  # points


def _header(ws: Worksheet, quote: Quote, settings: AppSettings, image: ImageResolver,
            last_col: str) -> int:  # fmt: skip
    lab = labels(quote.language)
    company = settings.company
    logo = image(company.logo_image_id) if company.logo_image_id else None
    if logo is not None:
        w, h = image_size(logo)
        img = XLImage(str(logo))
        img.height = 45
        img.width = int(45 * w / h) if h else 45
        img.anchor = "A1"
        ws.add_image(img)
    ws.merge_cells(f"A1:{last_col}1")
    ws["A1"] = company.name_en
    ws["A1"].font = Font(name="Calibri", bold=True, size=16)
    ws["A1"].alignment = CENTER
    ws.row_dimensions[1].height = 30
    ws.merge_cells(f"A2:{last_col}2")
    ws["A2"] = company.address_en
    ws["A2"].font = Font(name="Calibri", size=9)
    ws["A2"].alignment = CENTER
    customer = quote.customer_snapshot or {}
    top = [
        f"{lab['offer_no']}{quote.display_no}",
        f"{lab['to']} {customer.get('name') or ''}",
        f"{lab['contact']}{quote.contact or ''}",
    ]
    bottom = [
        f"{lab['offer_date']}{dmy(quote.offer_date)}",
        f"{lab['client_code']}{customer.get('code') or ''}",
        f"{lab['trade_term']} {quote.trade_term or ''}",
    ]
    if quote.contact_email or quote.validity_days:
        top.append(f"{lab['email']}{quote.contact_email or ''}")
        valid = f" {quote.validity_days} {lab['days']}" if quote.validity_days else ""
        bottom.append(f"{lab['valid']}{valid}")
    n_cols = ws.max_column if ws.max_column > 1 else len(LIST_WIDTHS)
    spans = _spans(n_cols, len(top))
    for row_no, values in ((3, top), (4, bottom)):
        for (first, last), value in zip(spans, values, strict=True):
            a, b = f"{get_column_letter(first)}{row_no}", f"{get_column_letter(last)}{row_no}"
            if first != last:
                _box_merge(ws, a, b)
            _set(ws, a, value, BOLD)
    return 6


def _spans(n_cols: int, parts: int) -> list[tuple[int, int]]:
    """Split columns 1..n into ``parts`` contiguous groups for the info grid."""
    if parts == 3 and n_cols == 6:
        return [(1, 2), (3, 3), (4, 6)]
    if parts == 3:
        return [(1, 2), (3, 5), (6, 8)]
    if n_cols == 6:
        return [(1, 2), (3, 3), (4, 5), (6, 6)]
    return [(1, 2), (3, 4), (5, 6), (7, 8)]


def _offer_section(ws: Worksheet, quote: Quote, section: QuoteSection, row: int) -> int:
    lab = labels(quote.language)
    titles = [lab["no"], lab["name"], lab["description"], lab["qty"], lab["unit_price"],
              lab["total_price"]]  # fmt: skip
    for col, title in enumerate(titles, start=1):
        _set(ws, f"{get_column_letter(col)}{row}", title, BOLD, CENTER, GREY)
    first = row + 1
    row += 1
    for number, line in zip(printed_numbers(section), section.lines, strict=True):
        red = Font(name="Calibri", size=10, color="FFE60000") if line.name_emphasis else NORMAL
        _set(ws, f"A{row}", number)
        _set(ws, f"B{row}", line.name, red)
        _set(ws, f"C{row}", _rich(line.description_lines or []), NORMAL, LEFT)
        _set(ws, f"D{row}", line.qty, number_format="#,##0")
        _set(
            ws,
            f"E{row}",
            line.unit_price,
            number_format=_money_format(line.unit_price, False, False),
        )
        _set(ws, f"F{row}", f"=D{row}*E{row}", NORMAL, RIGHT,
             number_format=_money_format(line.total, False, False))  # fmt: skip
        ws.row_dimensions[row].height = max(len(line.description_lines or []) * LINE_PT, 30)
        row += 1
    _box_merge(ws, f"A{row}", f"E{row}")
    _set(ws, f"A{row}", lab["total"], BOLD, CENTER, YELLOW)
    _set(ws, f"F{row}", f"=SUM(F{first}:F{row - 1})", BOLD, RIGHT, YELLOW,
         _money_format(section.subtotal, False, False))  # fmt: skip
    return row + 2


def _list_section(ws: Worksheet, quote: Quote, section: QuoteSection, row: int,
                  image: ImageResolver) -> int:  # fmt: skip
    lab = labels(quote.language)
    _box_merge(ws, f"A{row}", f"H{row}")
    _set(
        ws,
        f"A{row}",
        section.title or lab["quotation_list"],
        Font(name="Calibri", bold=True, size=14),
    )
    ws.row_dimensions[row].height = 24
    row += 1
    titles = [lab["list_no"], lab["description"], lab["model_no"], lab["material"], lab["photo"],
              lab["list_unit"], lab["quantity"], lab["list_total"]]  # fmt: skip
    for col, title in enumerate(titles, start=1):
        _set(ws, f"{get_column_letter(col)}{row}", title, BOLD)
    first = row + 1
    row += 1
    for number, line in zip(printed_numbers(section), section.lines, strict=True):
        _set(ws, f"A{row}", number, BOLD)
        _set(ws, f"B{row}", _rich(list_description(line)), NORMAL, LEFT)
        _set(ws, f"C{row}", line.model_no)
        _set(ws, f"D{row}", line.material)
        _set(ws, f"E{row}", None)
        photos = [p for p in (image(i) for i in (line.image_ids or [])[:2]) if p is not None]
        photo_h = _photos(ws, photos, 5, row) if photos else 0
        _set(
            ws,
            f"F{row}",
            line.unit_price,
            BOLD,
            number_format=_money_format(line.unit_price, True, False),
        )
        _set(ws, f"G{row}", line.qty, NORMAL, RIGHT, number_format="#,##0")
        _set(ws, f"H{row}", f"=F{row}*G{row}", number_format=_money_format(line.total, True, True))
        ws.row_dimensions[row].height = max(
            len(line.description_lines or []) * LINE_PT, photo_h, 30
        )
        row += 1
    _box_merge(ws, f"A{row}", f"G{row}")
    _set(ws, f"A{row}", lab["total"], BOLD, CENTER, YELLOW)
    _set(ws, f"H{row}", f"=SUM(H{first}:H{row - 1})", BOLD, CENTER, YELLOW,
         _money_format(section.subtotal, True, True))  # fmt: skip
    return row + 2


def quote_workbook(quote: Quote, settings: AppSettings, image: ImageResolver) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = quote.display_no[:31]
    has_list = any(s.layout == "list" for s in quote.sections)
    widths = LIST_WIDTHS if has_list else OFFER_WIDTHS
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.cell(row=1, column=len(widths))  # so max_column reflects the layout
    row = _header(ws, quote, settings, image, get_column_letter(len(widths)))
    for section in quote.sections:
        if section.layout == "list":
            row = _list_section(ws, quote, section, row, image)
        else:
            row = _offer_section(ws, quote, section, row)
    for text in (quote.payment_terms, quote.notes_footer):
        if text:
            ws.merge_cells(f"A{row}:{get_column_letter(len(widths))}{row}")
            ws[f"A{row}"] = text
            ws[f"A{row}"].alignment = LEFT
            ws.row_dimensions[row].height = max(LINE_PT * (len(text) // 120 + 1), LINE_PT)
            row += 1
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


PI_WIDTHS = (22, 16, 50, 26, 10, 13, 17)


def proforma_workbook(
    invoice: ProformaInvoice, settings: AppSettings, image: ImageResolver
) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = invoice.pi_no[:31]
    for index, width in enumerate(PI_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    company = settings.company
    customer = invoice.customer_snapshot or {}
    seller = invoice.seller_contact or {}
    ws["A1"] = company.name_en
    ws["A1"].font = Font(name="Arial", bold=True, size=13)
    ws["A2"] = company.address_en
    ws["A2"].font = Font(name="Arial", size=8)
    ws.merge_cells("A4:G4")
    ws["A4"] = "Proforma   Invoice"
    ws["A4"].font = Font(name="Arial", bold=True, size=15)
    ws["A4"].alignment = CENTER
    ws["A6"], ws["A6"].font = "Customer", BOLD
    ws["B7"], ws["B7"].font = customer.get("name") or "", Font(name="Arial", bold=True, size=11)
    ws["B9"], ws["B9"].font = f"Atten: {invoice.attention or ''}", BOLD
    info = [
        ("Inv No.:", invoice.pi_no), ("Date:", ymd_slash(invoice.issue_date)),
        ("P.O No.:", invoice.po_no or ""), ("Customer ID:", str(customer.get("code") or "")),
        ("Contact:", seller.get("name") or ""), ("Phone:", seller.get("phone") or ""),
        ("Currency:", invoice.currency),
    ]  # fmt: skip
    for offset, (label, value) in enumerate(info):
        ws[f"F{6 + offset}"], ws[f"F{6 + offset}"].font = label, BOLD
        ws[f"F{6 + offset}"].alignment = RIGHT
        ws[f"G{6 + offset}"] = value
    row = 14
    for col, title in enumerate(
        ("Model", "Name", "Description", "Photo", "Quantity", "Unit Price", "Net Amount"), start=1
    ):
        cell = ws.cell(row=row, column=col, value=title)
        cell.font, cell.alignment = BOLD, CENTER
        cell.border = Border(top=Side(style="medium"))
    first = row + 1
    row += 1
    foc = Decimal(invoice.foc_pct) if invoice.foc_pct else None
    for line in invoice.lines:
        lines = list(line.description_lines or [])
        if foc is not None:
            lines.append({"text": f"{foc.normalize():f}% FOC", "emphasis": False})
        ws[f"A{row}"], ws[f"A{row}"].font, ws[f"A{row}"].alignment = line.model_no, BOLD, CENTER
        ws[f"B{row}"], ws[f"B{row}"].font, ws[f"B{row}"].alignment = line.name, BOLD, CENTER
        ws[f"C{row}"], ws[f"C{row}"].alignment = _rich(lines), LEFT
        photos = [p for p in (image(i) for i in (line.image_ids or [])[:2]) if p is not None]
        photo_h = _photos(ws, photos, 4, row) if photos else 0
        ws[f"E{row}"], ws[f"E{row}"].alignment = line.qty, CENTER
        ws[f"F{row}"], ws[f"F{row}"].font = line.unit_price, BOLD
        ws[f"F{row}"].number_format = '"$"#,##0.00'
        ws[f"G{row}"], ws[f"G{row}"].font = f"=E{row}*F{row}", BOLD
        ws[f"G{row}"].number_format = '"$"#,##0.00'
        ws.row_dimensions[row].height = max(len(lines) * LINE_PT, photo_h, 30)
        row += 1
    for col in range(1, 8):
        ws.cell(row=row, column=col).border = Border(top=THIN)
    ws[f"B{row + 1}"], ws[f"B{row + 1}"].font = "Freight", BOLD
    if invoice.freight is not None:
        ws[f"G{row + 1}"] = invoice.freight
        ws[f"G{row + 1}"].number_format = '"$"#,##0.00'
    total_row = row + 3
    ws[f"F{total_row}"], ws[f"F{total_row}"].font = "Total", BOLD
    freight = f"+G{row + 1}" if invoice.freight is not None else ""
    ws[f"G{total_row}"] = f"=SUM(G{first}:G{row - 1}){freight}"
    ws[f"G{total_row}"].font = BOLD
    ws[f"G{total_row}"].number_format = '"$"#,##0.00'
    for letter in ("F", "G"):
        ws[f"{letter}{total_row}"].fill = PatternFill("solid", fgColor="FFE699")
    row = total_row + 2
    ws[f"B{row}"], ws[f"B{row}"].font = (
        "TERMS & CONDITION:",
        Font(name="Arial", bold=True, underline="single", size=9),
    )
    row += 1
    for number, term in enumerate(invoice.terms or [], start=1):
        ws[f"A{row}"], ws[f"A{row}"].alignment = number, RIGHT
        ws[f"B{row}"], ws[f"B{row}"].font = f"{term.get('title')}:", BOLD
        ws.merge_cells(f"C{row}:G{row}")
        ws[f"C{row}"], ws[f"C{row}"].alignment = term.get("text"), LEFT
        ws.row_dimensions[row].height = LINE_PT * (len(str(term.get("text"))) // 110 + 1)
        row += 1
    row += 1
    ws[f"B{row}"], ws[f"B{row}"].font = (
        "Our Banking Info:",
        Font(name="Arial", bold=True, underline="single", size=9),
    )
    for label, value in (
        ("Account Name:", company.bank_account_name), ("Account No:", company.bank_account_no),
        ("Bank Name:", company.bank_name), ("Swift Code :", company.bank_swift),
    ):  # fmt: skip
        row += 1
        ws[f"B{row}"] = label
        ws[f"C{row}"], ws[f"C{row}"].font = value, Font(name="Arial", bold=True, size=11)
    row += 2
    ws[f"B{row}"], ws[f"B{row}"].font = f"Seller:{company.name_en}", BOLD
    ws[f"E{row}"], ws[f"E{row}"].font = "Buyer:", BOLD
    ws[f"B{row + 2}"], ws[f"E{row + 2}"] = "Signed by:", "Signed by:"
    signature = image(company.signature_image_id) if company.signature_image_id else None
    if signature is not None:
        w, h = image_size(signature)
        img = XLImage(str(signature))
        img.height = 40
        img.width = int(40 * w / h) if h else 40
        img.anchor = f"C{row + 1}"
        ws.add_image(img)
    ws[f"B{row + 4}"] = f"Date: {ymd_slash(invoice.issue_date)}"
    ws[f"E{row + 4}"] = "Date:"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

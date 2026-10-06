"""The seller's internal version of a quote: every line with its cost, margin and profit.

Never meant for the customer: a red banner says so on the PDF and on the sheet. The customer
version stays ``quote_pdf`` / ``xlsx.quote_workbook``.
"""

from __future__ import annotations

import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.config import AppSettings
from app.models import Quote
from app.services.documents.common import (
    BLACK,
    FONT,
    RED,
    YELLOW,
    Cell,
    Column,
    Document,
    Run,
    Table,
    amount,
    dmy,
    labels,
    quantity,
    text_cell,
)
from app.services.documents.xlsx import BOLD, CENTER, GREY, LEFT, NORMAL, RIGHT, _set
from app.services.documents.xlsx import YELLOW as YELLOW_FILL
from app.services.quotes import CONDITIONS, Costing, PartCost

GREY_TEXT = (110, 110, 110)
SHARES = (0.05, 0.27, 0.07, 0.09, 0.08, 0.09, 0.12, 0.12, 0.11)
WIDTHS = (5, 40, 9, 12, 10, 12, 15, 15, 14)
STATUS_ZH = {"draft": "草稿", "sent": "已发送", "accepted": "已接受", "lost": "未成交",
             "expired": "已过期"}  # fmt: skip


def _titles(language: str) -> list[str]:
    lab = labels(language)
    return [lab["no"], lab["product"], lab["qty"], lab["unit_cost"], lab["margin"],
            lab["unit_price"], lab["cost_total"], lab["total_price"], lab["profit"]]  # fmt: skip


def _status(quote: Quote) -> str:
    return STATUS_ZH.get(quote.status, quote.status) if quote.language == "zh" else quote.status


def _pct(value: Decimal | None) -> str:
    return f"{Decimal(value):.2f}%" if value is not None else "—"


def _info(quote: Quote) -> list[tuple[str, str]]:
    lab = labels(quote.language)
    customer = quote.customer_snapshot or {}
    who = " ".join(str(v) for v in (customer.get("code"), customer.get("name")) if v)
    return [
        (lab["offer_no"], quote.display_no),
        (lab["customer"], who),
        (lab["offer_date"], dmy(quote.offer_date)),
        (lab["status"], _status(quote)),
    ]


def _summary(quote: Quote, cost: Costing) -> list[tuple[str, str]]:
    lab = labels(quote.language)
    rows = [(lab["total_cost"], amount(cost.cost, decimals="always", dollar=True))]
    if cost.discount:
        rows += [
            (lab["subtotal"], amount(Decimal(quote.subtotal), decimals="always", dollar=True)),
            (lab["discount"], amount(-cost.discount, decimals="always", dollar=True)),
        ]
    rows += [
        (lab["revenue"], amount(Decimal(quote.grand_total), decimals="always", dollar=True)),
        (lab["profit"].replace("($)", ""), amount(cost.profit, decimals="always", dollar=True)),
        (lab["margin_on_cost"], _pct(cost.margin_on_cost)),
    ]
    return rows


def _part_name(part: PartCost, language: str) -> str:
    c = part.component
    words = CONDITIONS["zh" if language == "zh" else "en"]
    name = str(c["name"])
    if c.get("condition") in words:
        name = f"{name} {words[c['condition']]}"
    return f"· {name}" + (f"  ×{c['qty']}" if c["qty"] > 1 else "")


def _part_cells(part: PartCost, units: int, language: str) -> list[Cell]:
    """A part under its combined product: grey, cost columns only."""
    c = part.component
    missing = part.cost_total is None
    cost = "—" if missing else amount(Decimal(c["cost_usd"]))
    grey = {"color": GREY_TEXT, "size": 7.5}
    return [
        text_cell(""),
        Cell(lines=[Run(_part_name(part, language), GREY_TEXT)], align="L", size=7.5),
        text_cell(quantity(c["qty"] * units), "R", **grey),
        text_cell(cost, "R", color=RED if missing else GREY_TEXT, size=7.5),
        text_cell(""),
        text_cell(""),
        text_cell(amount(part.cost_total, decimals="always") or "—", "R", **grey),
        text_cell(""),
        text_cell(""),
    ]


# --------------------------------------------------------------------- PDF


def render_costing(quote: Quote, cost: Costing, settings: AppSettings) -> bytes:
    lab = labels(quote.language)
    pdf = Document()
    pdf.set_title(f"{quote.display_no} internal")
    pdf.set_author(settings.company.name_en)
    pdf.add_page()
    width = pdf.content_width

    pdf.set_fill_color(*RED)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font(FONT, "B", 11)
    pdf.cell(width, 8, lab["internal"], align="C", fill=True, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    pdf.set_text_color(*BLACK)
    for label, value in _info(quote):
        pdf.set_font(FONT, "B", 9)
        pdf.cell(28, 5, label)
        pdf.set_font(FONT, "", 9)
        pdf.cell(width - 28, 5, value, new_x="LMARGIN", new_y="NEXT")

    columns = [Column(t, width * s) for t, s in zip(_titles(quote.language), SHARES, strict=True)]
    for index, section in enumerate(cost.sections, start=1):
        pdf.ln(4)
        table = Table(pdf, columns, x=pdf.l_margin, size=8, header_size=8)
        title = section.section.title or f"{lab['table']} {index}"
        table.row([Cell(lines=[Run(title, bold=True)], span=9, align="L", size=10)],
                  is_header=True)  # fmt: skip
        table.header()
        for item in section.lines:
            line = item.line
            name = [Run(line.name, RED if line.name_emphasis else BLACK)]
            if line.model_no and line.model_no != line.name:
                name.append(Run(line.model_no, GREY_TEXT))
            missing = item.cost_total is None
            table.row(
                [
                    text_cell(item.number),
                    Cell(lines=name, align="L"),
                    text_cell(quantity(line.qty), "R"),
                    text_cell(
                        amount(line.cost_usd) if not missing else "—",
                        "R",
                        color=RED if missing else BLACK,
                    ),  # fmt: skip
                    text_cell(_pct(item.margin), "R"),
                    text_cell(amount(line.unit_price), "R"),
                    text_cell(amount(item.cost_total, decimals="always") or "—", "R"),
                    text_cell(amount(line.total, decimals="always"), "R"),
                    text_cell(amount(item.profit, decimals="always") or "—", "R", bold=True),
                ]
            )
            for part in item.parts:
                table.row(_part_cells(part, line.qty, quote.language))
        table.row(
            [
                Cell(lines=[Run(lab["total"], bold=True)], span=6, fill=YELLOW),
                text_cell(amount(section.cost, decimals="always"), "R", bold=True, fill=YELLOW),
                text_cell(
                    amount(section.revenue, decimals="always"), "R", bold=True, fill=YELLOW
                ),  # fmt: skip
                text_cell(amount(section.profit, decimals="always"), "R", bold=True, fill=YELLOW),
            ]
        )

    rows = _summary(quote, cost)
    if pdf.get_y() + 6 + 6 * len(rows) + 10 > pdf.bottom:
        pdf.add_page()
    pdf.ln(4)
    pdf.c_margin = 1.5
    for label, value in rows:
        pdf.set_x(pdf.l_margin + width * 0.55)
        pdf.set_font(FONT, "B", 10)
        pdf.cell(width * 0.25, 6, label, border=1)
        pdf.cell(width * 0.2, 6, value, border=1, align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.c_margin = 0
    if cost.missing_cost:
        pdf.ln(2)
        pdf.set_font(FONT, "", 9)
        pdf.set_text_color(*RED)
        pdf.multi_cell(width, 4.5, lab["missing_cost"].format(n=cost.missing_cost), align="L")
        pdf.set_text_color(*BLACK)
    return bytes(pdf.output())


# ------------------------------------------------------------------- Excel


def costing_workbook(quote: Quote, cost: Costing) -> bytes:
    """The same table with formulas, so the seller can try other margins in Excel."""
    lab = labels(quote.language)
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = f"{quote.display_no[:22]} internal"
    for index, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(index)].width = w
    ws.merge_cells("A1:I1")
    _set(ws, "A1", lab["internal"], Font(name="Calibri", bold=True, size=12, color="FFFFFF"),
         CENTER, PatternFill("solid", fgColor="E60000"), border=False)  # fmt: skip
    ws.row_dimensions[1].height = 22
    row = 2
    for label, value in _info(quote):
        _set(ws, f"A{row}", label, BOLD, LEFT, border=False)
        ws.merge_cells(f"A{row}:B{row}")
        _set(ws, f"C{row}", value, NORMAL, LEFT, border=False)
        row += 1
    profit_cells: list[str] = []
    cost_cells: list[str] = []
    for index, section in enumerate(cost.sections, start=1):
        row += 1
        ws.merge_cells(f"A{row}:I{row}")
        _set(ws, f"A{row}", section.section.title or f"{lab['table']} {index}",
             Font(name="Calibri", bold=True, size=12), LEFT, border=False)  # fmt: skip
        row += 1
        for col, title in enumerate(_titles(quote.language), start=1):
            _set(ws, f"{get_column_letter(col)}{row}", title, BOLD, CENTER, GREY)
        row += 1
        mains: list[int] = []
        for item in section.lines:
            line = item.line
            name = line.name + (f"\n{line.model_no}" if line.model_no and line.model_no != line.name
                                else "")  # fmt: skip
            main = row
            mains.append(main)
            _set(ws, f"A{row}", item.number)
            _set(ws, f"B{row}", name, NORMAL, LEFT)
            _set(ws, f"C{row}", line.qty, NORMAL, RIGHT, number_format="#,##0")
            _set(ws, f"F{row}", line.unit_price, NORMAL, RIGHT, number_format="#,##0.00")
            _set(ws, f"H{row}", f"=C{row}*F{row}", NORMAL, RIGHT, number_format="#,##0.00")
            ws.row_dimensions[row].height = 28 if "\n" in name else 18
            if item.cost_total is None:
                for letter in "DEGI":
                    _set(ws, f"{letter}{row}", "—", Font(name="Calibri", size=10, color="E60000"))
            else:
                cost_value: object = line.cost_usd
                if item.parts:  # the parts below add up to the unit cost
                    last = row + len(item.parts)
                    cost_value = f"=SUM(G{row + 1}:G{last})/C{row}"
                _set(ws, f"D{row}", cost_value, NORMAL, RIGHT, number_format="#,##0.00##")
                _set(ws, f"E{row}", f"=IF(D{row}=0,\"\",F{row}/D{row}-1)", NORMAL, RIGHT,
                     number_format="0.00%")  # fmt: skip
                _set(ws, f"G{row}", f"=C{row}*D{row}", NORMAL, RIGHT, number_format="#,##0.00")
                _set(ws, f"I{row}", f"=H{row}-G{row}", BOLD, RIGHT, number_format="#,##0.00")
            row += 1
            for part in item.parts:
                c = part.component
                grey = Font(name="Calibri", size=9, color="6E6E6E")
                for letter in "AEFHI":
                    _set(ws, f"{letter}{row}", None)
                _set(ws, f"B{row}", _part_name(part, quote.language), grey, LEFT)
                _set(ws, f"C{row}", f"={c['qty']}*C{main}", grey, RIGHT, number_format="#,##0")
                if part.cost_total is None:
                    for letter in "DG":
                        _set(
                            ws, f"{letter}{row}", "—", Font(name="Calibri", size=9, color="E60000")
                        )
                else:
                    _set(ws, f"D{row}", Decimal(c["cost_usd"]), grey, RIGHT,
                         number_format="#,##0.00##")  # fmt: skip
                    _set(ws, f"G{row}", f"=C{row}*D{row}", grey, RIGHT, number_format="#,##0.00")
                row += 1
        ws.merge_cells(f"A{row}:F{row}")
        _set(ws, f"A{row}", lab["total"], BOLD, CENTER, YELLOW_FILL)
        for letter in "GHI":
            cells = "+".join(f"{letter}{r}" for r in mains) or "0"
            _set(ws, f"{letter}{row}", f"=SUM({cells})", BOLD, RIGHT, YELLOW_FILL, "#,##0.00")
        cost_cells.append(f"G{row}")
        profit_cells.append(f"I{row}")
        row += 1

    row += 1
    sums = "+".join(profit_cells) or "0"
    summary: list[tuple[str, str, object, str]] = [
        ("cost", lab["total_cost"], "=" + ("+".join(cost_cells) or "0"), "#,##0.00")
    ]
    if cost.discount:
        summary.append(("discount", lab["discount"], -cost.discount, "#,##0.00"))
    summary += [
        ("revenue", lab["revenue"], Decimal(quote.grand_total), "#,##0.00"),
        ("profit", lab["profit"], f"={sums}", "#,##0.00"),
        ("margin", lab["margin_on_cost"], "", "0.00%"),
    ]  # fmt: skip
    at = {key: row + index for index, (key, *_) in enumerate(summary)}
    for key, title, content, fmt in summary:
        if key == "margin":
            content = f'=IF(H{at["cost"]}=0,"",H{at["profit"]}/H{at["cost"]})'
        elif key == "profit" and "discount" in at:
            content = f"{content}+H{at['discount']}"  # the discount cell is negative
        ws.merge_cells(f"F{row}:G{row}")
        _set(ws, f"F{row}", title, BOLD, LEFT)
        _set(ws, f"H{row}", content, BOLD, RIGHT, number_format=fmt)
        row += 1
    if cost.missing_cost:
        ws.merge_cells(f"A{row + 1}:I{row + 1}")
        _set(ws, f"A{row + 1}", lab["missing_cost"].format(n=cost.missing_cost),
             Font(name="Calibri", size=10, color="E60000"), LEFT, border=False)  # fmt: skip
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
